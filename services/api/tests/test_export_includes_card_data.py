"""An export that leaves the instalments behind is not an export.

Everything else in the file can be rebuilt from the ledger. An instalment plan
cannot: it is entered by hand, once, and nothing else in the export implies it.
The same goes for the two billing days on a card - without them a restored
account is an ordinary liability and every due date it carried is gone.

So this is the one omission that loses data rather than convenience, and it is
what somebody discovers only after they have exported everything and started
again.
"""
import io
import json
import uuid
from datetime import datetime, timezone

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.security import create_access_token, hash_password
from app.db.database import get_db
from app.models.models import Account, Base, Emi, User
from main import app

pytestmark = pytest.mark.asyncio


@pytest_asyncio.fixture
async def api():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(bind=engine, class_=AsyncSession, expire_on_commit=False)

    async def _db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = _db

    user_id, card_id, plain_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    async with factory() as session:
        session.add(User(id=user_id, email="export@example.com",
                         password_hash=hash_password("x"), display_name="Export"))
        session.add(Account(id=card_id, user_id=user_id, name="HDFC Regalia",
                            account_type="liability", currency="INR",
                            opening_balance_minor=0,
                            statement_day=18, due_day=8, credit_limit_minor=2_000_000))
        session.add(Account(id=plain_id, user_id=user_id, name="SBI Savings",
                            account_type="asset", currency="INR", opening_balance_minor=0))
        await session.commit()
        session.add(Emi(user_id=user_id, account_id=card_id, name="iPhone 16",
                        monthly_minor=650000, months=12, currency="INR",
                        started_at=datetime(2026, 3, 10, tzinfo=timezone.utc)))
        session.add(Emi(user_id=user_id, account_id=None, name="Fridge",
                        monthly_minor=300000, months=6, currency="INR", is_active=False,
                        started_at=datetime(2025, 11, 1, tzinfo=timezone.utc)))
        await session.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        client.headers["Authorization"] = f"Bearer {create_access_token({'sub': str(user_id)})}"
        yield client

    app.dependency_overrides.clear()


class TestTheSpreadsheet:
    async def test_there_is_a_sheet_for_instalments(self, api):
        res = await api.get("/api/profile/export.xlsx")
        assert res.status_code == 200, res.text[:200]
        wb = load_workbook(io.BytesIO(res.content))
        assert "Instalments" in wb.sheetnames

    async def test_every_plan_is_in_it_finished_ones_included(self, api):
        """A closed plan is history, and history is what an export is for."""
        wb = load_workbook(io.BytesIO((await api.get("/api/profile/export.xlsx")).content))
        names = [row[0] for row in wb["Instalments"].iter_rows(min_row=2, values_only=True)]
        assert sorted(n for n in names if n) == ["Fridge", "iPhone 16"]

    async def test_a_plan_carries_what_it_takes_to_rebuild_it(self, api):
        wb = load_workbook(io.BytesIO((await api.get("/api/profile/export.xlsx")).content))
        ws = wb["Instalments"]
        headers = [c.value for c in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "iPhone 16")
        got = dict(zip(headers, row))

        assert got["Instalment (INR)"] == 6500.00
        assert got["Instalments"] == 12
        assert str(got["First instalment"]).startswith("2026-03-10")
        # The figure the plan exists to make visible.
        assert got["Total (INR)"] == 78000.00
        assert got["Charged to"] == "HDFC Regalia"
        assert got["Status"] == "Running"

    async def test_a_closed_plan_says_so(self, api):
        wb = load_workbook(io.BytesIO((await api.get("/api/profile/export.xlsx")).content))
        ws = wb["Instalments"]
        headers = [c.value for c in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "Fridge")
        got = dict(zip(headers, row))
        assert got["Status"] == "Closed"
        # Charged to no account, which is legitimate. openpyxl reads an empty
        # cell back as None - what matters is that it is BLANK and not the
        # string "None", which is what a careless str() would have put there.
        assert got["Charged to"] in (None, "")

    async def test_the_accounts_sheet_carries_the_billing_days(self, api):
        """Without these the card is just a liability, and the cycle is gone."""
        wb = load_workbook(io.BytesIO((await api.get("/api/profile/export.xlsx")).content))
        ws = wb["Accounts"]
        headers = [c.value for c in ws[1]]
        rows = {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)}

        card = rows["HDFC Regalia"]
        assert card["Statement day"] == 18
        assert card["Payment due day"] == 8
        assert card["Credit limit (INR)"] == 20000.00

        # An ordinary account has none of them, and shows blanks rather than
        # zeros - a zero would read as a real day of the month, and "None" as
        # a value somebody had entered.
        plain = rows["SBI Savings"]
        # `in (None, "")` already excludes the string "None" - which is the
        # thing being guarded against, and what a careless str() would leave
        # in the cell for somebody to read as a real entry.
        for field in ("Statement day", "Payment due day", "Credit limit (INR)"):
            assert plain[field] in (None, ""), f"{field} came out as {plain[field]!r}"

    async def test_no_cell_anywhere_is_a_formula(self, api):
        """The export-wide rule has to hold on the new sheet too.

        A plan is named by the user, so it is user data like any other, and
        openpyxl stores a leading "=" as a live formula.
        """
        wb = load_workbook(io.BytesIO((await api.get("/api/profile/export.xlsx")).content))
        formulas = [
            f"{sheet.title}!{cell.coordinate}"
            for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row
            if cell.data_type == "f"
        ]
        assert formulas == [], f"these would execute on open: {formulas}"


class TestTheJsonExport:
    async def test_it_carries_the_plans(self, api):
        res = await api.get("/api/profile/export")
        assert res.status_code == 200, res.text[:200]
        data = res.json()
        assert "emis" in data, "the plans are missing from the JSON export"
        by_name = {e["name"]: e for e in data["emis"]}
        assert set(by_name) == {"iPhone 16", "Fridge"}
        assert by_name["iPhone 16"]["monthly_minor"] == 650000
        assert by_name["iPhone 16"]["months"] == 12
        assert by_name["Fridge"]["is_active"] is False

    async def test_it_carries_the_billing_days(self, api):
        data = (await api.get("/api/profile/export")).json()
        by_name = {a["name"]: a for a in data["accounts"]}
        assert by_name["HDFC Regalia"]["statement_day"] == 18
        assert by_name["HDFC Regalia"]["due_day"] == 8
        assert by_name["SBI Savings"]["statement_day"] is None

    async def test_it_is_still_valid_json(self, api):
        raw = (await api.get("/api/profile/export")).content
        assert isinstance(json.loads(raw), dict)
