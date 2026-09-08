import uuid
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, Query, Response
from sqlalchemy import select, and_, or_, desc
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import get_current_user
from app.core.ratelimit import IMPORT_BY_ACCOUNT, SHEET_BY_ACCOUNT, enforce
from app.db.database import get_db
from app.models.models import User, Account, Category, Transaction, SavingsGoal
from app.schemas.schemas import (
    TransactionCreate, TransactionUpdate, TransactionResponse,
    ImportRequest, ImportResult, ImportRejection,
    SheetRequest, SheetResult,
)

router = APIRouter(prefix="/transactions", tags=["Transactions"])

@router.get("", response_model=List[TransactionResponse])
async def list_transactions(
    account_id: Optional[uuid.UUID] = Query(None),
    category_id: Optional[uuid.UUID] = Query(None),
    transaction_type: Optional[str] = Query(None),
    start_date: Optional[datetime] = Query(None),
    end_date: Optional[datetime] = Query(None),
    # The dashboard needs the five most recent rows. Without a limit it had to
    # download every transaction the user has ever recorded to show them.
    limit: Optional[int] = Query(None, ge=1, le=500),
    offset: int = Query(0, ge=0),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """Lists transactions for current user with optional filtering by account, category, date, or type."""
    filters = [Transaction.user_id == current_user.id]
    if account_id:
        filters.append(Transaction.account_id == account_id)
    if category_id:
        filters.append(Transaction.category_id == category_id)
    if transaction_type:
        filters.append(Transaction.transaction_type == transaction_type.lower())
    if start_date:
        filters.append(Transaction.transaction_date >= start_date)
    if end_date:
        filters.append(Transaction.transaction_date <= end_date)

    stmt = select(Transaction).where(and_(*filters)).order_by(
        desc(Transaction.transaction_date), desc(Transaction.id)
    )
    if offset:
        stmt = stmt.offset(offset)
    if limit is not None:
        stmt = stmt.limit(limit)
    res = await db.execute(stmt)
    return res.scalars().all()


@router.post("", response_model=TransactionResponse, status_code=status.HTTP_201_CREATED)
async def create_transaction(
    payload: TransactionCreate,
    response: Response,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Creates a new transaction idempotently.
    Repeated requests with the same client_mutation_id return the existing record cleanly.
    """
    # Read once, up front - see the recovery path at the end of this function.
    # After a rollback every ORM object is expired, and reading an attribute
    # off one then issues a blocking reload that an async session cannot serve.
    user_id = current_user.id

    # 1. Idempotency Check
    existing_stmt = select(Transaction).where(Transaction.client_mutation_id == payload.client_mutation_id)
    existing_res = await db.execute(existing_stmt)
    existing_tx = existing_res.scalar_one_or_none()
    if existing_tx:
        if existing_tx.user_id != current_user.id:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied to mutation record.")
        response.status_code = status.HTTP_200_OK
        return existing_tx

    # 2. Validation
    if payload.amount_minor <= 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Transaction amount_minor must be greater than zero."
        )

    if payload.transaction_type.lower() not in ["income", "expense", "transfer"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="transaction_type must be 'income', 'expense', or 'transfer'."
        )

    # Account ownership check
    acc_stmt = select(Account).where(and_(Account.id == payload.account_id, Account.user_id == current_user.id))
    acc_res = await db.execute(acc_stmt)
    account = acc_res.scalar_one_or_none()
    if not account:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Account not found or access denied.")

    # Transfer target account / savings goal validation
    if payload.transaction_type.lower() == "transfer":
        if not payload.to_account_id and not payload.savings_goal_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Transfers require either a to_account_id or a savings_goal_id."
            )
        if payload.to_account_id:
            if payload.account_id == payload.to_account_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Source and target transfer accounts cannot be the same."
                )
            to_acc_stmt = select(Account).where(and_(Account.id == payload.to_account_id, Account.user_id == current_user.id))
            to_acc_res = await db.execute(to_acc_stmt)
            if not to_acc_res.scalar_one_or_none():
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail="Target transfer account not found or access denied."
                )

    # Category validation
    if payload.category_id:
        cat_stmt = select(Category).where(
            and_(
                Category.id == payload.category_id,
                or_(Category.user_id == current_user.id, Category.user_id == None)
            )
        )
        cat_res = await db.execute(cat_stmt)
        if not cat_res.scalar_one_or_none():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Category not found or access denied.")

    # Savings goal validation
    if payload.savings_goal_id:
        goal_stmt = select(SavingsGoal).where(and_(SavingsGoal.id == payload.savings_goal_id, SavingsGoal.user_id == current_user.id))
        goal_res = await db.execute(goal_stmt)
        if not goal_res.scalar_one_or_none():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Savings goal not found or access denied.")

    # 3. Create Transaction record
    tx = Transaction(
        client_mutation_id=payload.client_mutation_id,
        user_id=current_user.id,
        account_id=payload.account_id,
        to_account_id=payload.to_account_id if payload.transaction_type.lower() == "transfer" else None,
        category_id=payload.category_id,
        savings_goal_id=payload.savings_goal_id,
        transaction_type=payload.transaction_type.lower(),
        amount_minor=payload.amount_minor,
        currency=payload.currency or current_user.currency,
        description=payload.description,
        transaction_date=payload.transaction_date,
        device_id=payload.device_id,
        sync_status=payload.sync_status or "synced"
    )
    db.add(tx)

    try:
        await db.commit()
        await db.refresh(tx)
        return tx
    except IntegrityError:
        await db.rollback()
        # Fallback concurrency race check
        fallback_res = await db.execute(
            select(Transaction).where(
                and_(
                    Transaction.client_mutation_id == payload.client_mutation_id,
                    Transaction.user_id == user_id
                )
            )
        )
        existing_race_tx = fallback_res.scalar_one_or_none()
        if existing_race_tx:
            return existing_race_tx
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Transaction mutation conflict.")


@router.get("/{transaction_id}", response_model=TransactionResponse)
async def get_transaction(
    transaction_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """Retrieves a single transaction by ID."""
    stmt = select(Transaction).where(and_(Transaction.id == transaction_id, Transaction.user_id == current_user.id))
    res = await db.execute(stmt)
    tx = res.scalar_one_or_none()
    if not tx:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Transaction not found.")
    return tx


@router.patch("/{transaction_id}", response_model=TransactionResponse)
async def update_transaction(
    transaction_id: uuid.UUID,
    payload: TransactionUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """Updates non-mutation fields of a transaction with ownership and concurrency validation."""
    stmt = select(Transaction).where(and_(Transaction.id == transaction_id, Transaction.user_id == current_user.id))
    res = await db.execute(stmt)
    tx = res.scalar_one_or_none()

    if not tx:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Transaction not found.")

    if payload.expected_version is not None and tx.version != payload.expected_version:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Transaction update conflict: version mismatch."
        )

    if payload.amount_minor is not None:
        if payload.amount_minor <= 0:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Amount must be greater than zero.")
        tx.amount_minor = payload.amount_minor
    if payload.description is not None:
        tx.description = payload.description
    if payload.transaction_date is not None:
        tx.transaction_date = payload.transaction_date

    if payload.category_id is not None:
        cat_stmt = select(Category).where(
            and_(
                Category.id == payload.category_id,
                or_(Category.user_id == current_user.id, Category.user_id == None)
            )
        )
        cat_res = await db.execute(cat_stmt)
        if not cat_res.scalar_one_or_none():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Category not found or access denied.")
        tx.category_id = payload.category_id

    if payload.savings_goal_id is not None:
        goal_stmt = select(SavingsGoal).where(and_(SavingsGoal.id == payload.savings_goal_id, SavingsGoal.user_id == current_user.id))
        goal_res = await db.execute(goal_stmt)
        if not goal_res.scalar_one_or_none():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Savings goal not found or access denied.")
        tx.savings_goal_id = payload.savings_goal_id

    tx.version += 1
    await db.commit()
    await db.refresh(tx)
    return tx


@router.delete("/{transaction_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_transaction(
    transaction_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """Deletes a transaction from the ledger."""
    stmt = select(Transaction).where(and_(Transaction.id == transaction_id, Transaction.user_id == current_user.id))
    res = await db.execute(stmt)
    tx = res.scalar_one_or_none()

    if not tx:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Transaction not found.")

    await db.delete(tx)
    await db.commit()
    return None


@router.post("/import", response_model=ImportResult)
async def import_transactions(
    payload: ImportRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Takes a chunk of a bank statement.

    Exists because the alternative is one HTTP request per line. A year of
    statements is several hundred rows, and on a phone connection that is
    minutes of round trips and a half-finished import whenever the signal
    drops.

    The property that matters is that running it TWICE is safe. People import
    the same file again, or next month's export which overlaps last month's,
    and the honest answer to that is "nothing happened" rather than a second
    copy of everything. Each row carries an id derived from the payment
    itself, so a repeat is recognised by identity rather than by guessing at
    similar-looking rows.

    Duplicates are counted, not rejected: an overlapping export is the normal
    case, not an error, and reporting it as one would teach people to ignore
    the result.
    """
    user_id = current_user.id
    # Bounded per account. Generous - a decade of history is fifty chunks -
    # but it stops a runaway loop from writing until the disk gives out.
    enforce(IMPORT_BY_ACCOUNT, str(user_id))

    acc_res = await db.execute(
        select(Account).where(and_(Account.id == payload.account_id, Account.user_id == user_id))
    )
    if not acc_res.scalar_one_or_none():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Account not found or access denied.",
        )

    incoming_ids = [row.client_mutation_id for row in payload.rows]

    # Tried twice at most. The second pass exists for one case: another import
    # of the same file committing between the lookup below and this one's
    # commit. The first attempt then dies on the unique constraint and the
    # whole chunk rolls back - including rows that were genuinely new. Rolling
    # back and reporting them all as duplicates, which is what this used to do,
    # loses those rows AND tells the user they are safely stored.
    for attempt in range(2):
        try:
            return await _apply_import(db, user_id, payload, incoming_ids)
        except IntegrityError:
            await db.rollback()
            if attempt == 1:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Another import was running. Nothing was lost - try this file again.",
                )
    # Unreachable; the loop either returns or raises.
    raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Import could not complete.")


async def _apply_import(db, user_id, payload, incoming_ids) -> ImportResult:
    """One attempt. Re-reads what exists, so a retry sees the winner's rows."""
    # Every id in the chunk, looked up once. Row-by-row checks would be one
    # query per line - the same round-trip problem this endpoint exists to fix,
    # moved from the network to the database.
    seen_res = await db.execute(
        select(Transaction.client_mutation_id, Transaction.user_id)
        .where(Transaction.client_mutation_id.in_(incoming_ids))
    )
    seen = {mutation_id: owner for mutation_id, owner in seen_res.all()}

    created = 0
    duplicates = 0
    rejected: list[ImportRejection] = []
    # A file can repeat a line, and two identical rows in one chunk would both
    # pass the check above and collide on the constraint at commit.
    in_this_chunk: set = set()

    for index, row in enumerate(payload.rows):
        if row.client_mutation_id in seen:
            # Somebody else's id is not this user's duplicate to absorb.
            if seen[row.client_mutation_id] != user_id:
                # Deliberately vague. Saying WHY would confirm that this id
                # exists in somebody else's ledger, and the ids are derived
                # from the payment - so a crafted row could be used to ask
                # "does this person have a payment of X to Y on date Z?" and
                # get a definitive answer. The ids are now scoped per user as
                # well, which is the real fix; this is the second lock.
                rejected.append(ImportRejection(index=index, reason="Could not be added."))
            else:
                duplicates += 1
            continue

        if row.client_mutation_id in in_this_chunk:
            duplicates += 1
            continue

        if row.amount_minor <= 0:
            rejected.append(ImportRejection(index=index, reason="Amount must be greater than zero."))
            continue

        kind = row.transaction_type.lower()
        # Deliberately not "transfer": a statement line cannot say which of the
        # user's own accounts the money went to, and a transfer without a
        # destination is one the ledger cannot balance.
        if kind not in ("income", "expense"):
            rejected.append(
                ImportRejection(index=index, reason="Only income and expense can be imported.")
            )
            continue

        db.add(
            Transaction(
                client_mutation_id=row.client_mutation_id,
                user_id=user_id,
                account_id=payload.account_id,
                transaction_type=kind,
                amount_minor=row.amount_minor,
                currency=row.currency,
                description=row.description,
                transaction_date=row.transaction_date,
                device_id="statement-import",
            )
        )
        in_this_chunk.add(row.client_mutation_id)
        created += 1

    # Deliberately NOT caught here. A constraint failure means the picture
    # this pass built is stale, and the caller retries against a fresh read
    # rather than guessing at what landed.
    await db.commit()
    return ImportResult(created=created, duplicates=duplicates, rejected=rejected)


#: A statement of a few hundred kilobytes, with room to spare. Bounded because
#: this is decoded into memory on a shared instance.
MAX_SHEET_BYTES = 6 * 1024 * 1024

#: The same bound, applied to what the archive says it EXPANDS to.
#:
#: The one above is on compressed bytes, and an .xlsx is a ZIP - a few hundred
#: kilobytes of repetitive XML decompresses to gigabytes. A limit on the
#: compressed size alone is not a limit at all against a file built to exploit
#: exactly that. Sixty megabytes of sheet XML is far past any real statement.
MAX_SHEET_UNCOMPRESSED = 60 * 1024 * 1024


@router.post("/import/sheet", response_model=SheetResult)
async def read_sheet(
    payload: SheetRequest,
    current_user: User = Depends(get_current_user),
):
    """Turns an .xlsx statement into CSV text.

    Banks push Excel harder than CSV, and until now the answer to an .xlsx was
    "look for the CSV download instead" - which for several banks does not
    exist. openpyxl is already a dependency here, for the export in the other
    direction, so the conversion costs nothing new.

    It converts and stops. Reading the columns, the dates and the amounts stays
    in the client's parser, which has tests against real bank layouts; a second
    implementation here would be a second set of rules to keep in step, and the
    one place they drifted apart would be a silent, wrong import.

    Dates are written back as ISO. A spreadsheet stores a real date rather than
    text, so this is the one route where the day/month ambiguity that plagues
    CSV can be settled properly instead of inferred.
    """
    import base64
    import csv as csv_module
    import io as io_module
    import zipfile

    # A decompress plus an XML parse costs far more than storing rows.
    enforce(SHEET_BY_ACCOUNT, str(current_user.id))

    try:
        raw = base64.b64decode(payload.content_base64, validate=True)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That file could not be read.",
        )

    if len(raw) > MAX_SHEET_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That file is too large. Export a shorter period and try again.",
        )

    # An .xlsx is a ZIP, and the cap above is on the COMPRESSED bytes. A few
    # hundred kilobytes of highly repetitive XML decompresses to gigabytes -
    # a zip bomb - and openpyxl would happily start expanding it before
    # anything noticed. So the archive is inspected before it is parsed: the
    # declared uncompressed sizes are read from the directory, which costs
    # nothing and does not decompress a single byte.
    try:
        with zipfile.ZipFile(io_module.BytesIO(raw)) as archive:
            declared = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That did not look like a spreadsheet. A .xlsx or .csv export works.",
        )

    if declared > MAX_SHEET_UNCOMPRESSED:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That file expands to more than this can read. Export a shorter period.",
        )

    try:
        from openpyxl import load_workbook

        # read_only keeps a long statement from being built as objects; data_only
        # takes the cached value of a formula rather than the formula itself,
        # which is what a balance column usually is.
        workbook = load_workbook(
            io_module.BytesIO(raw), read_only=True, data_only=True
        )
        sheet = workbook[workbook.sheetnames[0]]

        out = io_module.StringIO()
        writer = csv_module.writer(out, lineterminator="\n")
        written = 0
        for row in sheet.iter_rows(values_only=True):
            cells = []
            for value in row:
                if value is None:
                    cells.append("")
                elif isinstance(value, datetime):
                    # ISO, so the client reads it as year-month-day and never
                    # has to guess which of the first two numbers is the day.
                    cells.append(value.date().isoformat())
                else:
                    cells.append(str(value))
            # Trailing empty columns are an artefact of the sheet's used range.
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            writer.writerow(cells)
            written += 1
            if written > 20_000:      # far past what one import will take
                break

        workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That did not look like a spreadsheet. A .xlsx or .csv export works.",
        )

    return SheetResult(csv=out.getvalue(), sheet_name=str(sheet.title), rows=written)
